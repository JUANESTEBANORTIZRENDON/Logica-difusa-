"""Exportaciones trazables de simulaciones e informe docente."""

from __future__ import annotations

import csv
import io
import json
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from docx import Document

from model.config import MODEL_VERSION

BOGOTA = ZoneInfo("America/Bogota")


def timestamp() -> str:
    return datetime.now(BOGOTA).isoformat(timespec="seconds")


def serializable_result(result: dict) -> dict:
    return {
        "version_modelo": result["version_modelo"],
        "fecha_zona_horaria": timestamp(),
        "entrada": result["entrada"],
        "grados_entrada": result["grados_entrada"],
        "reglas": result["reglas"],
        "resultado": {
            "indice": result["indice"],
            "nivel": result["nivel"],
            "area": result["area"],
            "grados_salida": result["grados_salida"],
        },
    }


def result_json(result: dict) -> bytes:
    return json.dumps(serializable_result(result), ensure_ascii=False, indent=2).encode("utf-8")


def result_csv(result: dict) -> bytes:
    data = serializable_result(result)
    stream = io.StringIO(newline="")
    fields = ["version_modelo", "fecha_zona_horaria", "peso_kg", "edad_anios", "regla", "peso_conjunto", "edad_conjunto", "consecuente", "peso_regla", "alpha", "indice", "nivel"]
    writer = csv.DictWriter(stream, fieldnames=fields)
    writer.writeheader()
    for rule in data["reglas"]:
        writer.writerow({
            "version_modelo": data["version_modelo"],
            "fecha_zona_horaria": data["fecha_zona_horaria"],
            "peso_kg": data["entrada"]["peso"],
            "edad_anios": data["entrada"]["edad"],
            "regla": rule["id"],
            "peso_conjunto": rule["peso_conjunto"],
            "edad_conjunto": rule["edad_conjunto"],
            "consecuente": rule["consecuente"],
            "peso_regla": rule["peso_regla"],
            "alpha": rule["alpha"],
            "indice": data["resultado"]["indice"],
            "nivel": data["resultado"]["nivel"],
        })
    return stream.getvalue().encode("utf-8-sig")


def _document_bytes(document: Document) -> bytes:
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def simulation_docx(result: dict) -> bytes:
    data = serializable_result(result)
    doc = Document()
    doc.add_heading("Informe de simulación difusa", 0)
    doc.add_paragraph("Uso exclusivamente académico. No es probabilidad clínica ni diagnóstico.")
    doc.add_paragraph(f"Modelo: {MODEL_VERSION} | Fecha: {data['fecha_zona_horaria']}")
    doc.add_heading("Entradas y resultado", level=1)
    doc.add_paragraph(f"Peso: {data['entrada']['peso']} kg\nEdad: {data['entrada']['edad']} años\nÍndice: {data['resultado']['indice']:.4f} puntos\nEtiqueta: {data['resultado']['nivel']}")
    doc.add_heading("Grados de entrada", level=1)
    doc.add_paragraph(json.dumps(data["grados_entrada"], ensure_ascii=False, indent=2))
    doc.add_heading("Traza de las nueve reglas", level=1)
    table = doc.add_table(rows=1, cols=5)
    for cell, value in zip(table.rows[0].cells, ["Regla", "Peso", "Edad", "Consecuente", "Alfa"]):
        cell.text = value
    for rule in data["reglas"]:
        cells = table.add_row().cells
        for cell, value in zip(cells, [rule["id"], rule["peso_conjunto"], rule["edad_conjunto"], rule["consecuente"], f"{rule['alpha']:.6f}"]):
            cell.text = str(value)
    return _document_bytes(doc)


def teacher_docx(markdown_path: Path | None = None) -> bytes:
    source = markdown_path or Path(__file__).resolve().parents[1] / "docs" / "contenido.md"
    root = source.resolve().parents[1]
    doc = Document()
    for raw_line in source.read_text(encoding="utf-8").splitlines():
        line = raw_line.strip()
        if line.startswith("# "):
            doc.add_heading(line[2:], 0)
        elif line.startswith("## "):
            doc.add_heading(line[3:], 1)
        elif line.startswith("### "):
            doc.add_heading(line[4:], 2)
        elif line.startswith("- "):
            doc.add_paragraph(line[2:], style="List Bullet")
        elif line and not line.startswith("```"):
            doc.add_paragraph(line)
    doc.add_heading("Código ejecutable completo del motor", level=1)
    for relative in ("model/config.py", "model/membership.py", "model/inference.py"):
        doc.add_heading(relative, level=2)
        doc.add_paragraph((root / relative).read_text(encoding="utf-8"))
    return _document_bytes(doc)
